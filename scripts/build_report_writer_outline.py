from __future__ import annotations

from pathlib import Path

from build_initial_owner_brief import add_body, add_bullet, add_heading, add_table, set_run_font
from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUTPUT = Path("docs/submission/初赛材料/给报告负责人的初赛文档框架.docx")


def paragraph(doc: DocxDocument, lead: str, text: str) -> None:
    add_body(doc, f"{lead}：{text}", bold_lead=f"{lead}：")


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_style = doc.styles["Title"]
    title_style.font.name = "Microsoft YaHei"
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.font.underline = False
    p_pr = title_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(15)
    set_run_font(title.add_run("智析 Data Agent 初赛报告写作框架"), size=20, bold=True)

    paragraph(doc, "收件人", "报告主笔及负责提供素材的四位成员。")
    paragraph(
        doc,
        "用途",
        (
            "请按本框架扩写 S1 方案概要与 S2 详细解决方案，"
            "并从同一内容提炼 PPT、演示视频和企业专项材料。"
        ),
    )
    paragraph(
        doc,
        "提交原则",
        "角色与贡献写真实；功能状态以代码、实机截图和验收记录为准；图像生成仅用于概念示意，不能冒充系统运行截图。",
    )

    add_heading(doc, "一 团队角色与责任人")
    add_body(
        doc,
        "徐介翰是已确定的组长，建议在报告中担任项目经理兼技术经理。其余四人的项目岗位按现有任务安排提出，定稿前请按实际贡献核对；学号仍需按报名表补齐。PPT、视频和文稿属于交付任务，不作为项目岗位名称。",
    )
    add_table(
        doc,
        ["成员", "项目岗位", "项目职责", "提供给报告主笔的事实"],
        [
            [
                "徐介翰［学号待填］",
                "项目经理兼技术经理",
                "项目计划、版本决策、系统研发、技术架构和功能验收",
                "项目定位、S2B 过程记录、S2C 技术事实、真实截图和演示环境",
            ],
            [
                "孙竞杰［学号待填］",
                "产品经理",
                "用户需求、功能范围、使用流程、交互验收和产品价值",
                "S2A 用户场景、服务流程、使用手册和产品界面说明",
            ],
            [
                "马星宇［学号待填］",
                "业务分析经理",
                "制造业务场景、指标口径、服务模式和成本假设",
                "业务知识、典型案例、S2D 成本与可行性资料",
            ],
            [
                "陈庚琦［学号待填］",
                "交互设计经理",
                "用户操作路径、分析结果呈现、信息层级与展示一致性",
                "交互流程图、页面图示、产品效果与图源说明",
            ],
            [
                "代洛瑞［学号待填］",
                "测试与交付经理",
                "演示场景验证、问题记录、发布检查和交付质量",
                "稳定演示步骤、测试结果、视频素材与交付记录",
            ],
        ],
        [2.5, 3.6, 4.4, 5.6],
    )
    paragraph(
        doc,
        "责任划分",
        (
            "S2A 由产品经理提供需求与流程，S2B 由项目经理提供组织与过程，"
            "S2C 由技术经理提供架构与实现，S2D 由业务分析经理提供成本与可行性。"
            "负责文稿的同学负责文字整合，各岗位核对自己提供的事实。"
        ),
    )

    add_heading(doc, "二 S1 方案概要")
    add_body(
        doc,
        (
            "控制在 800 字以内，建议按四段写：企业数据底座已有数据却难以自助分析；"
            "智析采用业务语义加受控 Agent 工具链；展示自然语言提问、智能执行、"
            "表格图表和报告输出；说明多轮追问、证据追溯及企业价值。"
            "具体功能状态必须与最终演示版本一致。"
        ),
    )
    paragraph(
        doc,
        "主笔需向组长索取",
        "最终项目名称、演示版本功能清单、两组能稳定复现的分析案例、系统主页与结果页真实截图。",
    )

    add_heading(doc, "三 S2 详细解决方案")
    add_heading(doc, "S2A 目标与服务模型", 2)
    for item in [
        "目标问题：业务人员不懂表名、字段、SQL 和指标公式，数据工程师人工响应慢，口径难统一。",
        "目标客户与使用者：已有数据底座的制造企业；质量、生产、设备、库存、管理和数据管理员。",
        (
            "服务流程：接入只读数据源，扫描表与字段，确认业务语义，"
            "提出自然语言问题，Agent 选工具执行，返回结果并追问。"
        ),
        "产品价值：缩短数据定位与分析准备时间，保留统一指标口径，让分析过程可复核。不要写未经测量的效率提升百分比。",
        "配图建议：用户旅程图一张，实机的对话与结果截图两张。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "S2B 组织管理与业务分析", 2)
    for item in [
        "组织结构：使用第一节的五人岗位表，写清项目经理、技术经理、产品经理、业务分析经理、交互设计经理和测试与交付经理及兼任关系。",
        (
            "过程管理：记录需求拆解、原型设计、数据准备、迭代开发、"
            "测试验收、文档制作和合稿；以 Git 提交、验收文件和会议记录作证据。"
        ),
        "三天冲刺安排：第一天统一版本和报告框架，第二天冻结演示版并补齐素材，第三天合稿、录制、复核和上传。",
        "风险与对策：模型不稳定、演示环境中断、数据口径不一致、图文状态不一致、视频体积超限；每项写责任人与备用方案。",
        "配图建议：组织分工图、三天甘特图、系统迭代或测试记录截图。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "S2C 技术路线及实现方案", 2)
    for item in [
        "总体架构：交互层、Agent 编排层、业务语义层、查询分析层、数据治理层和平台基础层。",
        "核心闭环：用户问题、意图与上下文、目录和语义绑定、工具选择、安全执行、结果验证、自然语言解释与继续追问。",
        "数据与知识：表、字段、类型、说明、样例值、关系、业务对象、指标、规则和分析主题；强调面向未知表的扫描与人工确认映射。",
        (
            "执行与安全：展示生成 SQL 或脚本的路径；已实现的受控 SQL 与只读权限如实描述；"
            "尚未完成的代码沙箱和六类建模任务列为后续计划。"
        ),
        "结果呈现：实机表格、图表、证据定位、报告；逐项标注已实现、开发中或规划，并写出验收证据。",
        (
            "配图建议：六层架构图、Agent 工具调用流程图、真实语义管理页、"
            "真实对话页、真实图表与查询依据页。"
        ),
    ]:
        add_bullet(doc, item)
    paragraph(
        doc,
        "技术经理提供",
        "最终技术栈、工具清单、架构图校稿、数据表和指标清单、实机截图、关键测试结果及功能状态。",
    )

    add_heading(doc, "S2D 成本模型及可行性分析", 2)
    for item in [
        "成本口径：按模型调用、应用服务器、数据库、对象存储、运维和客户实施列成本项；填入实际测算的单价、用量和区间。",
        "交付模式：可讨论企业私有化部署与托管服务两种路径，明确数据留存、权限和运维责任。",
        "技术可行性：用已运行的端到端流程、真实数据库接入、可复现样例和测试记录支撑。",
        "商业可行性：说明目标客户场景、采购动机、试点方式和后续扩展，避免无依据市场规模或收入预测。",
        "配图建议：成本测算表、部署模式对比图、试点路线图。",
    ]:
        add_bullet(doc, item)
    paragraph(
        doc,
        "仍待填写",
        (
            "模型 API 价格与估算用量、服务器规格、月度运行费用、实施周期和试点客户假设。"
            "所有数字标注日期、来源及假设。"
        ),
    )

    doc.add_page_break()  # type: ignore[no-untyped-call]
    add_heading(doc, "四 其他材料与报告主笔的交接清单")
    add_table(
        doc,
        ["材料", "报告主笔需准备或核对", "素材来源"],
        [
            [
                "S3 项目简介 PPT",
                "与报告一致的价值主张、核心方案、实机效果和十分钟讲述顺序",
                "产品经理与交互设计经理",
            ],
            [
                "S4A 可运行原型",
                "登录方式、启动方式、样例数据和稳定演示流程",
                "项目经理与测试交付经理",
            ],
            [
                "S4B 演示视频",
                "不超过十分钟，压缩后不超过 150MB；与最终版实机一致",
                "测试与交付经理",
            ],
            [
                "S5 团队完成过程",
                "不超过五分钟；写实记录角色、分工、排期、协作、测试与复盘",
                "项目经理汇总全员过程证据",
            ],
            [
                "企业专项材料",
                "产品使用手册、产品交互演示、分工与过程、业务知识说明、数据资源说明",
                "文稿同学整合；各项目岗位提供事实",
            ],
        ],
        [3.0, 7.1, 6.0],
    )
    paragraph(
        doc,
        "图像标注",
        (
            "实机截图写版本、页面、问题和结果；架构图写为设计示意；"
            "AI 生成图写为概念示意。报告中不把概念图当成已运行功能。"
        ),
    )
    paragraph(
        doc,
        "交稿前核对",
        (
            "S1 字数、S2A 至 S2D 齐全、五人角色与报名信息一致、技术状态与演示一致、"
            "视频时长和体积合规、全部文件按官方方式命名。"
        ),
    )

    doc.save(str(OUTPUT))
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
