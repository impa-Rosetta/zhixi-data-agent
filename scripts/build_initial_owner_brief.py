from __future__ import annotations

from pathlib import Path
from typing import cast

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table, _Cell, _Row
from docx.text.paragraph import Paragraph
from docx.text.run import Run

OUTPUT = Path("docs/submission/初赛材料/智析DataAgent初赛项目负责人工作稿.docx")


def set_cell_shading(cell: _Cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(
    cell: _Cell, top: int = 120, start: int = 120, bottom: int = 120, end: int = 120
) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row: _Row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row: _Row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_run_font(
    run: Run, name: str = "Microsoft YaHei", size: float = 10.5, bold: bool = False
) -> None:
    run.font.name = name
    fonts = run._element.get_or_add_rPr().rFonts
    assert fonts is not None
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_paragraph_spacing(
    paragraph: Paragraph, before: float = 0, after: float = 6, line: float = 1.35
) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_body(doc: DocxDocument, text: str, *, bold_lead: str | None = None) -> Paragraph:
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Normal"]
    set_paragraph_spacing(paragraph)
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet(doc: DocxDocument, text: str, level: int = 0) -> Paragraph:
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_paragraph_spacing(paragraph, after=3)
    set_run_font(paragraph.add_run(text))
    return paragraph


def add_number(doc: DocxDocument, text: str) -> Paragraph:
    paragraph = doc.add_paragraph(style="List Number")
    set_paragraph_spacing(paragraph, after=4)
    set_run_font(paragraph.add_run(text))
    return paragraph


def add_heading(doc: DocxDocument, text: str, level: int = 1) -> Paragraph:
    paragraph = cast(Paragraph, doc.add_heading(text, level=level))
    set_paragraph_spacing(paragraph, before=12 if level == 1 else 8, after=6, line=1.15)
    for run in paragraph.runs:
        set_run_font(run, size=15 if level == 1 else 12, bold=True)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_table(
    doc: DocxDocument, headers: list[str], rows: list[list[str]], widths: list[float]
) -> Table:
    table = cast(Table, doc.add_table(rows=1, cols=len(headers)))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(paragraph, after=0, line=1.1)
        run = paragraph.add_run(value)
        set_run_font(run, size=9.5, bold=True)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        row = table.add_row()  # type: ignore[no-untyped-call]
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, "F4F7FB")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            )
            set_paragraph_spacing(paragraph, after=0, line=1.15)
            set_run_font(paragraph.add_run(value), size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title, before=44, after=12, line=1.0)
    set_run_font(title.add_run("智析 Data Agent 初赛项目负责人工作稿"), size=23, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(subtitle, after=24)
    run = subtitle.add_run("统一项目口径  材料总框架  三天交付计划")
    set_run_font(run, size=12)
    run.font.color.rgb = RGBColor(80, 80, 80)

    info = add_table(
        doc,
        ["项目", "内容"],
        [
            ["赛题", "A07 企业数据底座智能问析 Agent 系统"],
            ["项目名称", "智析 企业数据底座智能问析 Agent 系统"],
            ["负责人", "待填写 姓名及学号"],
            ["团队", "共五人 成员顺序须与报名表一致"],
            ["交付周期", "三天初赛材料冲刺"],
            ["最终截止时间", "待填写"],
        ],
        [3.2, 12.5],
    )
    info.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_body(
        doc,
        "本文档由项目负责人维护，用于统一系统完善、报告、PPT和视频的叙事口径。团队成员应以本稿确定的项目定位、功能状态、演示数据和术语为准，再分别完成自己的交付材料。",
    )
    add_body(
        doc,
        "当前结论：项目已经具备企业数据接入、业务语义、自然语言问析、多轮对话、安全查询、图表和证据追溯主链路。三天冲刺阶段不再追求功能数量，而应优先保证核心演示稳定、材料真实一致、重点优势清晰可见。",
        bold_lead="当前结论：",
    )

    add_heading(doc, "一 项目统一定位")
    add_heading(doc, "一句话定位", 2)
    add_body(
        doc,
        (
            "智析 Data Agent 是面向企业数据底座的可信智能问析系统。"
            "它把企业数据表、字段、关系和指标口径组织为可执行的业务语义，"
            "由 Agent 根据用户自然语言问题选择合适的数据目录、查询、统计、图表和报告工具，"
            "形成从提问到可信结果再到持续追问的完整闭环。"
        ),
    )
    add_heading(doc, "目标用户", 2)
    for item in [
        "质量人员：查询不良率、一次通过率、缺陷排行和质量趋势。",
        "生产人员：分析产量、工序、产线和工单执行情况。",
        "设备人员：分析停机、报警、运行效率及其与质量指标的关系。",
        "管理人员：通过自然语言快速获得周报、月报和专题分析结果。",
        "数据管理员：维护数据源、目录、字段画像、业务指标和语义映射。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "核心价值", 2)
    add_body(
        doc,
        "传统流程需要业务人员提出需求，再由数据工程师定位数据、编写SQL、制作图表和解释结果。智析把重复过程沉淀为可复用的Agent工具链，降低业务人员使用企业数据底座的门槛，同时通过权限、安全验证和证据链控制大模型的不确定性。",
    )

    add_heading(doc, "二 需要统一突出的项目优势")
    advantages = [
        [
            "1",
            "受控 Agent 闭环",
            "系统不是只生成文字，而是完成理解、规划、选工具、执行、验证和呈现。",
        ],
        [
            "2",
            "语义驱动而非裸 Text to SQL",
            "模型负责理解表达，确定性语义模型负责指标公式、维度和SQL生成，减少口径漂移。",
        ],
        ["3", "可信和可追溯", "结果保留查询、产物、验证和Evidence引用，用户可以定位数字依据。"],
        [
            "4",
            "连续多轮交互",
            "用户可以在同一会话追问、补充条件和切换主题，Agent使用自然语言澄清。",
        ],
        [
            "5",
            "未知数据资源适配",
            "系统先扫描表、字段、关系和画像，再由管理员确认物理字段与业务语义的映射。",
        ],
        [
            "6",
            "企业级安全边界",
            "工作空间隔离、角色权限、只读执行、SQL AST门禁、脱敏、审计和任务恢复共同约束执行。",
        ],
        [
            "7",
            "可产品化交付",
            "FastAPI、React、PostgreSQL、Redis、Celery和MinIO构成可容器部署的完整产品，而非单页演示。",
        ],
    ]
    add_table(doc, ["序号", "优势", "统一表述"], advantages, [1.2, 3.8, 10.7])

    add_heading(doc, "三 初赛解决方案总框架")
    add_body(
        doc,
        (
            "详细解决方案严格按照官方 S2A 至 S2D 组织。"
            "所有组员制作的报告、PPT和视频必须能够映射回以下四部分，"
            "避免技术内容很多但没有形成完整商业解决方案。"
        ),
    )
    s2_rows = [
        [
            "S2A",
            "目标与服务模型",
            "问题背景、目标用户、服务流程、产品价值、典型场景和与传统方案的差异。",
        ],
        [
            "S2B",
            "组织管理与业务分析",
            "五人分工、研发流程、里程碑、质量管理、风险管理、目标客户和推广方式。",
        ],
        [
            "S2C",
            "技术路线及实现方案",
            "数据接入、元数据、语义模型、Agent、工具调用、安全查询、多轮交互、图表报告、权限审计和部署。",
        ],
        [
            "S2D",
            "成本模型及可行性",
            "模型API、服务器、数据库、对象存储和运维成本；私有化与SaaS模式；技术与商业可行性。",
        ],
    ]
    add_table(doc, ["部分", "主题", "必须包含"], s2_rows, [1.7, 4.0, 10.0])

    add_heading(doc, "四 当前系统状态和表述边界")
    status_rows = [
        [
            "可以写成已实现",
            (
                "登录与工作空间、成员权限、PostgreSQL和MySQL安全接入、元数据目录与画像、"
                "制造质量语义指标、安全查询、DeepSeek Agent、多轮对话、自然澄清、"
                "SSE恢复、表格图表和Evidence。"
            ),
        ],
        [
            "已完成后端核心 仍需前端收口",
            "可信报告持久化、编排、创建API、异步生成Markdown HTML PDF和MinIO发布。",
        ],
        [
            "不得写成已完整实现",
            "六类机器学习模型、完整隔离Python沙箱、报告预览下载前端、评测中心、OpenTelemetry、生产级性能和备份恢复。",
        ],
        [
            "材料统一规则",
            "已实现能力使用真实截图和测试依据；开发中能力注明当前阶段；规划能力只写路线图，不用AI效果图伪装成真实功能。",
        ],
    ]
    add_table(doc, ["状态", "统一口径"], status_rows, [4.0, 11.7])

    add_heading(doc, "五 三天内系统完善范围")
    add_body(
        doc,
        "系统工作采用P0、P1、P2三级管理。负责人只承诺完成P0，P1在不影响稳定性的前提下推进，P2不进入本轮开发。",
    )
    priority_rows = [
        [
            "P0 必须完成",
            "一键启动稳定；固定账号和样例数据可用；登录、会话、Agent回答、多轮追问、表格、图表和Evidence主链路可重复演示；页面无明显错误；敏感信息不出现在截图和视频。",
        ],
        [
            "P1 尽量完成",
            "可信报告的预览、下载和重试闭环；核心演示问题的回答质量；移动端关键页面；统一错误提示。",
        ],
        [
            "P2 暂不扩展",
            "完整六模型平台、评测中心、复杂可观测性、全量生产设备库存场景和与演示无关的重构。",
        ],
    ]
    add_table(doc, ["优先级", "范围"], priority_rows, [3.4, 12.3])
    add_body(
        doc,
        "版本冻结规则：第二天中午冻结演示版本。冻结后只修复会导致无法启动、无法登录、核心分析失败、页面严重错位或敏感信息泄漏的问题。所有截图和录屏均来自冻结版本。",
        bold_lead="版本冻结规则：",
    )

    add_heading(doc, "六 统一演示主线")
    demo_steps = [
        "管理员通过一键启动进入系统并登录演示账号。",
        "展示已接入的数据源、表、字段、关系、画像和制造质量语义指标。",
        "用户创建会话并提问 最近三个月不良率趋势。",
        "Agent识别指标和时间范围，选择目录、语义查询、安全验证和图表工具。",
        "系统返回自然语言结论、数据表、折线图、可信等级和Evidence。",
        "用户在同一会话追问 哪个月最高 与上个月相比变化多少。",
        "展示Agent继承上下文，并根据已有结果选择合适工具完成分析。",
        "展示权限、安全SQL、查询依据或技术详情，证明结果不是模型编造。",
        "如报告功能已稳定，生成质量分析报告并展示预览或下载；如果尚未完成，只在路线图中说明。",
        "最后总结企业价值、可推广能力和下一阶段规划。",
    ]
    for item in demo_steps:
        add_number(doc, item)

    add_heading(doc, "备用演示问题")
    for item in [
        "请分析各工序的良率。",
        "找出最近一个月不良数量最高的产品。",
        "统计每条产线最近七天的产量趋势。",
        "设备停机时间和不良率是否相关。",
        "生成一份本周质量分析结果。",
        "你好 你能帮我做什么。",
        "如果问题缺少时间或指标口径，验证Agent是否能自然追问。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "七 初赛材料总框架")
    material_rows = [
        [
            "S1",
            "方案概要",
            "不超过800字，说明目标问题、解决思路、创新点、成果与价值。",
            "报告负责人",
        ],
        [
            "S2",
            "详细解决方案",
            "按S2A至S2D组织，加入团队责任、成本和可行性。",
            "负责人定框架 报告负责人完善",
        ],
        [
            "S3",
            "项目简介PPT",
            "控制在10分钟，强调真实产品、核心优势、演示结果和企业价值。",
            "PPT内容与视觉负责人",
        ],
        [
            "S4A",
            "可运行系统",
            "冻结演示版本、一键启动、固定数据、演示账号和操作说明。",
            "项目负责人",
        ],
        ["S4B", "演示视频", "不超过10分钟且小于150MB，展示完整Agent闭环。", "视频负责人"],
        [
            "S5",
            "团队完成过程",
            "PPT或视频，不超过5分钟，展示分工、计划、协作、学习、创新和执行。",
            "PPT内容负责人",
        ],
        ["企业1", "产品使用手册", "功能架构、角色、流程图、主要操作和典型案例。", "报告负责人"],
        ["企业2", "产品交互演示", "突出操作过程，可复用S4B素材但应单独导出。", "视频负责人"],
        [
            "企业3",
            "分工及过程文档",
            "五人角色、排期、决策、Git记录、测试和验收证据。",
            "报告负责人",
        ],
        ["企业4", "业务知识说明", "业务对象、指标、规则、主题、公式和分析口径。", "报告负责人"],
        [
            "企业5",
            "数据资源说明",
            "表、字段、关系、样例值、数据字典、来源和脱敏说明。",
            "报告负责人",
        ],
    ]
    add_table(doc, ["编号", "材料", "内容要求", "负责人"], material_rows, [1.4, 3.0, 8.2, 3.1])

    add_heading(doc, "八 五人协作接口")
    role_rows = [
        [
            "项目负责人",
            "冻结产品、确定框架、提供真实素材、审核事实边界、最终提交。",
            "功能清单、统一口径、截图目录、演示账号、最终批准。",
        ],
        [
            "报告负责人",
            "把框架扩写成正式材料，统一文字、图表、图片和引用。",
            "S1、S2、手册、知识说明、数据说明、过程文档。",
        ],
        ["PPT内容负责人", "把长文稿转化为10分钟叙事和讲解稿。", "逐页文案、讲解词、S5内容。"],
        [
            "PPT视觉负责人",
            "维护唯一正式PPT文件，统一版式、配色、截图和动画。",
            "S3正式PPT、配图、打包目录检查。",
        ],
        ["视频负责人", "分镜、录屏、剪辑、字幕、配音、压缩和隐私检查。", "S4B和产品交互演示视频。"],
    ]
    add_table(doc, ["角色", "责任边界", "主要输出"], role_rows, [3.0, 7.0, 5.7])
    add_body(
        doc,
        "协作规则：每项材料只能有一名最终文件负责人。其他成员通过文字稿、截图或批注协作，不同时维护多个正式版本。文件名必须含版本号和更新时间，最终版统一放入提交目录。",
        bold_lead="协作规则：",
    )

    add_heading(doc, "九 三天执行计划")
    plan_rows = [
        ["第一天 上午", "负责人发布统一口径、演示主线、材料目录和真实功能清单；全员确认分工。"],
        [
            "第一天 下午",
            "系统修复P0问题；报告完成主文档初稿；PPT完成逐页框架；视频完成分镜和试录。",
        ],
        ["第一天 晚上", "进行第一次联合审查，解决项目名称、术语、功能状态和视觉风格不一致。"],
        ["第二天 上午", "继续修复演示链路；报告、PPT和视频按真实系统补充素材。"],
        ["第二天 中午", "冻结演示版本、账号和数据。冻结后统一截图和正式录屏。"],
        ["第二天 下午", "完成正式截图、主报告二稿、PPT整稿和视频粗剪；拆分企业专项材料。"],
        ["第三天 上午", "交叉校对内容、演练PPT、检查视频字幕和功能表述，修复阻断问题。"],
        ["第三天 下午", "导出PDF、压缩视频、测试PPT、测试系统、核对命名并生成最终RAR。"],
        ["截止前两小时", "停止改稿，换一台电脑解压和复查最终包，完成上传。"],
    ]
    add_table(doc, ["时间", "必须完成"], plan_rows, [3.2, 12.5])

    add_heading(doc, "十 最终审核清单")
    checklist = [
        "项目名称、赛题编号、学校、成员、学号和报名顺序全部一致。",
        "S1不超过800字，S3和S4B均能在10分钟内完成，S5不超过5分钟。",
        "视频文件小于150MB，常见播放器可以正常打开。",
        "所有功能截图来自真实系统，AI图片只用于概念表达并检查中文文字。",
        "报告、PPT和视频对功能状态的描述一致，不把规划能力写成已完成。",
        "截图和录屏不出现API Key、密码、数据库连接串、真实个人信息或内部错误详情。",
        "演示账号、样例数据、一键启动和核心问题在冻结版本上至少完整走通两次。",
        "Word、PDF、PPT和视频均在另一台电脑上打开检查。",
        "外层RAR文件名严格按照报名成员顺序命名。",
        "最终RAR解压后目录清晰、无临时文件、无重复版本、无损坏文件。",
    ]
    for item in checklist:
        add_bullet(doc, "□ " + item)

    add_heading(doc, "十一 负责人需要补充的信息")
    fields = [
        ["团队信息", "学校、指导老师、五名成员姓名、学号、报名顺序和联系方式。"],
        ["时间信息", "正式截止时间、上传平台、单文件限制、压缩包内部命名规则。"],
        ["分工信息", "每位成员的真实贡献、负责材料和可证明的过程记录。"],
        ["演示信息", "最终演示账号、登录密码、演示数据版本、核心问题和预期结果。"],
        ["成本信息", "拟采用的模型价格、服务器规格和私有化或SaaS商业模式假设。"],
        ["视觉信息", "最终Logo、主色、报告与PPT模板选择。"],
    ]
    add_table(doc, ["待补项目", "需要填写"], fields, [3.4, 12.3])

    add_heading(doc, "附录 建议素材目录")
    for item in [
        "01 登录与工作空间页面",
        "02 数据源列表与连接配置",
        "03 数据目录 表字段关系与画像",
        "04 业务指标和语义模型",
        "05 Agent自然语言提问",
        "06 Agent计划与工具执行过程",
        "07 表格与图表结果",
        "08 多轮追问与自然澄清",
        "09 Evidence与查询验证",
        "10 报告生成与下载 如果已完成",
        "11 团队过程 Git提交 测试和验收记录",
    ]:
        add_bullet(doc, item)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("智析 Data Agent 初赛项目负责人工作稿"), size=8)
    footer.runs[0].font.color.rgb = RGBColor(100, 100, 100)

    doc.save(str(OUTPUT))
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
