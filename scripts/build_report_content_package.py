from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from build_initial_owner_brief import add_body, add_bullet, add_heading, add_table, set_run_font


SOURCE = Path("docs/submission/初赛材料/给报告主笔的完整项目内容资料包.md")
OUTPUT = SOURCE.with_suffix(".docx")


def tune_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")


def add_markdown_table(doc: Document, lines: list[str]) -> None:
    records = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    if len(records) < 3:
        return
    headers, separator, *rows = records
    if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in separator):
        raise ValueError(f"Invalid table separator: {separator}")
    if len(headers) == 3:
        widths = [3.0, 7.4, 5.8]
    elif len(headers) == 2:
        widths = [4.0, 12.2]
    else:
        widths = [16.2 / len(headers)] * len(headers)
    table = add_table(doc, headers, rows, widths)
    tune_table_borders(table)


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(paragraph.add_run("智析 Data Agent 报告主笔资料包  ·  "), size=8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def build() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_style = doc.styles["Title"]
    title_style.font.name = "Microsoft YaHei"
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.font.underline = False
    p_pr = title_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)
    add_page_number(section)

    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            title = doc.add_paragraph(style="Title")
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title.paragraph_format.space_before = Pt(20)
            title.paragraph_format.space_after = Pt(14)
            set_run_font(title.add_run(line[2:]), size=20, bold=True)
            index += 1
            continue
        if line.startswith("## "):
            add_heading(doc, line[3:], level=1)
            index += 1
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:], level=2)
            index += 1
            continue
        if line.startswith("| "):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_markdown_table(doc, table_lines)
            continue
        if line.startswith("- "):
            add_bullet(doc, line[2:])
            index += 1
            continue
        if re.match(r"\d+\.\s", line):
            add_body(doc, line)
            index += 1
            continue

        paragraph_lines = [line.rstrip(" ")]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith(("#", "|", "- ")) or re.match(r"\d+\.\s", candidate):
                break
            paragraph_lines.append(candidate)
            index += 1
        add_body(doc, " ".join(paragraph_lines))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
